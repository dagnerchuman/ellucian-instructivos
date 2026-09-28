# -*- coding: utf-8 -*-
"""Planilla de revisión de la migración R2 para Centros Empresariales (USS)."""
import os
from datetime import date

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import PageSetupProperties

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, '..', 'entregables', 'PLANILLA DE REVISIÓN - MIGRACIÓN R2 - CENTROS EMPRESARIALES.xlsx')

F = 'Arial'
PURPLE, PURPLE_L, GREEN_L, AMBER_L, RED_L, GRAY_L, INPUT = '7030A0', 'F1E7FA', 'E2F0D9', 'FFF2CC', 'FDE2E2', 'F2F2F2', 'FFF9DB'
thin = Side(style='thin', color='D9D2E3')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD_FONT = Font(name=F, bold=True, color='FFFFFF', size=10)
HEAD_FILL = PatternFill('solid', fgColor=PURPLE)
BODY = Font(name=F, size=10)
BOLD = Font(name=F, size=10, bold=True)
WRAP = Alignment(wrap_text=True, vertical='top')
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)

REV_ROWS, ISS_ROWS = 400, 200

CENTROS = ['Idiomas', 'Computación', 'Emprendimiento']
PLANTILLAS = ['01 Persona', '01a Datos adicionales', '01b Escuela de procedencia', '01b Contacto de emergencia',
              '02 Estudiantes', '03 Historia académica', '04 Docentes', '05 Saldos', '06 Egresados']
CRITERIOS = [  # criterio, descripción, umbral
    ('Integridad', 'Campos obligatorios completos según las reglas funcionales definidas.', 0.99),
    ('Exactitud', 'Valores en Banner Student coinciden con la fuente de origen o evidencia validada.', 0.98),
    ('Consistencia', 'Datos relacionados coherentes entre módulos, tablas y registros (estudiante, programa, periodo, inscripción).', 0.98),
    ('Validez', 'Formatos, códigos, dominios, fechas y valores cumplen las reglas y catálogos configurados en Banner.', 0.99),
    ('Unicidad', 'Ausencia de registros duplicados de personas, estudiantes, inscripciones u otras entidades únicas.', 0.995),
    ('Integridad referencial', 'Las relaciones entre registros se mantienen y los totales concilian con la fuente de origen.', 0.99),
]
RESULTADOS = ['Conforme', 'Observado', 'No conforme']
SEVERIDAD = ['Alta', 'Media', 'Baja']
ESTADOS = ['Abierto', 'En corrección', 'Resuelto', 'Cerrado']
APLICA = ['Sí', 'No', 'Por confirmar']

# Tiempos por alumno (Ellucian, diapositiva 9): plantilla, qué validar, página, minutos, etapa del cronograma, aplica a CE
TIEMPOS = [
    ('01 Persona', 'Identificación, nombres, datos biográficos, documento, direcciones, teléfonos y correos.', 'SPAIDEN', 5, 'Personas', 'Sí'),
    ('01 Persona', 'Familiares o personas relacionadas con el estudiante.', 'SOAFOLK', 3, 'Personas', 'Sí'),
    ('01 Persona', 'Información médica del estudiante.', 'GOAMEDI', 2, 'Personas', 'Sí'),
    ('01 Persona', 'Retenciones del estudiante.', 'SOAHOLD', 2, 'Personas', 'Sí'),
    ('01a Datos adicionales', 'Otros tipos de documentos de identidad.', 'GVAADID', 2, 'Documento de identidad', 'Sí'),
    ('01b Escuela de procedencia', 'Escuela/colegio de procedencia, institución, tipo y atributos.', 'SOAPCOL', 3, 'Escuela de procedencia', 'Por confirmar'),
    ('01b Contacto de emergencia', 'Contactos de emergencia, relación, nombres y teléfonos.', 'SPAEMRG', 2, 'Contacto de emergencia', 'Sí'),
    ('02 Estudiantes', 'Programa, nivel, plan, campus, periodo de ingreso, estados, versión de currículo y programas concurrentes.', 'SGASTDN', 5, 'Estudiantes', 'Sí'),
    ('02 Estudiantes', 'Cohortes y atributos.', 'SGASADD', 2, 'Estudiantes', 'Sí'),
    ('03 Historia académica', 'Cursos, periodos, créditos, calificaciones y consistencia del historial.', 'SHACRSE + SHATCKN', 10, 'Historia académica', 'Sí'),
    ('03 Historia académica', 'Promedio por periodo y acumulado.', 'SHATERM + SHAINST', 6, 'Historia académica', 'Sí'),
    ('03 Historia académica', 'Ejecución de la evaluación de cumplimiento (CAPP).', 'SMARQCM', 3, 'Historia académica', 'Sí'),
    ('03 Historia académica', 'Revisión de resultados de cumplimiento; plan de equivalencia.', 'SMICRLT', 10, 'Historia académica', 'Sí'),
    ('03 Historia académica', 'Ejecución de la proyección de cursos.', 'SFPPROJ', 3, 'Historia académica', 'Sí'),
    ('03 Historia académica', 'Resultados de la proyección: prerrequisitos, equivalencias, versiones de malla.', 'SFAPROJ', 10, 'Historia académica', 'Sí'),
    ('04 Docentes', 'Datos personales, estado, tipo de instructor, relación institucional y grados (tiempo por docente).', 'SIAINST', 5, 'Docentes', 'Sí'),
    ('05 Saldos', 'Saldos, conceptos y valores frente al sistema anterior; saldo cero y pendiente.', 'TSAAREV', 5, 'Saldos', 'Sí'),
    ('06 Egresados', 'Graduación, programa, plan, fecha de grado y consistencia con la trayectoria.', 'SGASTDN + SHADEGR', 5, 'Egresados', 'Por confirmar'),
    ('07 Tesis', 'Registro de tesis y relación con el proceso de grado.', 'SHAQPNO', 3, 'Tesis', 'No'),
]
PAGINAS = sorted({p for t in TIEMPOS for p in t[2].split(' + ')} | {'GOAMTCH', 'SFAREGS', 'SIAASGN'})

# Cronograma Migración R2 (Ellucian, diapositiva 12): etapa, ambiente, carga ini, carga fin, rev ini, rev fin, muestra (tipo)
CRONO = [
    ('Personas', 'PROD', date(2026, 9, 7), date(2026, 9, 21), date(2026, 9, 24), date(2026, 9, 27), 'part'),
    ('Documento de identidad', 'PROD', date(2026, 9, 22), date(2026, 9, 25), date(2026, 9, 25), date(2026, 9, 27), 'part'),
    ('Contacto de emergencia', 'PROD', date(2026, 9, 22), date(2026, 9, 25), date(2026, 9, 25), date(2026, 9, 27), 'part'),
    ('Docentes', 'PROD', date(2026, 9, 28), date(2026, 9, 30), date(2026, 9, 30), date(2026, 10, 2), 'doc'),
    ('Carga LD01 con equivalencias', 'PROD', date(2026, 9, 30), date(2026, 10, 2), None, None, None),
    ('Matriz de seguridad', 'PROD', date(2026, 9, 30), date(2026, 10, 2), None, None, None),
    ('Clonación a TEST', '—', date(2026, 10, 5), date(2026, 10, 9), None, None, None),
    ('Estudiantes', 'TEST', date(2026, 10, 12), date(2026, 10, 16), date(2026, 10, 16), date(2026, 10, 21), 'part'),
    ('Historia académica', 'TEST', date(2026, 10, 19), date(2026, 10, 28), date(2026, 10, 28), date(2026, 10, 30), 'part'),
    ('Egresados', 'TEST', date(2026, 10, 29), date(2026, 11, 2), date(2026, 11, 2), date(2026, 11, 4), 'egr'),
    ('Saldos', 'TEST', date(2026, 11, 3), date(2026, 11, 6), date(2026, 11, 6), date(2026, 11, 13), 'part'),
    ('Tesis', 'TEST', date(2026, 11, 9), date(2026, 11, 10), date(2026, 11, 10), date(2026, 11, 13), None),
    ('Escuela de procedencia', 'TEST', date(2026, 11, 11), date(2026, 11, 13), date(2026, 11, 13), date(2026, 11, 16), 'part'),
    ('Pruebas integrales (USS)', 'TEST', None, None, date(2026, 11, 16), date(2026, 12, 26), None),
]


def style_header(ws, row, headers, widths):
    for i, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row=row, column=i, value=h)
        c.font, c.fill, c.alignment, c.border = HEAD_FONT, HEAD_FILL, CENTER, BORDER
        if w:
            ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 30


def title(ws, text, sub=None):
    ws['A1'] = text
    ws['A1'].font = Font(name=F, size=14, bold=True, color=PURPLE)
    if sub:
        ws['A2'] = sub
        ws['A2'].font = Font(name=F, size=9, italic=True, color='595959')


def body_cell(c, fill=None, align=WRAP, font=BODY, fmt=None):
    c.font, c.alignment, c.border = font, align, BORDER
    if fill:
        c.fill = PatternFill('solid', fgColor=fill)
    if fmt:
        c.number_format = fmt


def add_list_validation(ws, rng, ref, prompt):
    dv = DataValidation(type='list', formula1=ref, allow_blank=True, showDropDown=False)
    dv.promptTitle, dv.prompt = 'Seleccione', prompt
    dv.error, dv.errorTitle, dv.showErrorMessage = 'Elija un valor de la lista.', 'Valor no válido', True
    ws.add_data_validation(dv)
    dv.add(rng)


def result_colors(ws, rng):
    ws.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"Conforme"'], fill=PatternFill('solid', fgColor=GREEN_L)))
    ws.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"Observado"'], fill=PatternFill('solid', fgColor=AMBER_L)))
    ws.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"No conforme"'], fill=PatternFill('solid', fgColor=RED_L)))


def build():
    wb = Workbook()

    # ------------------------------------------------------------ Listas (oculta)
    ls = wb.active
    ls.title = 'Listas'
    lists = {'A': ('Centro', CENTROS), 'B': ('Plantilla', PLANTILLAS), 'C': ('Página', PAGINAS),
             'D': ('Criterio', [c[0] for c in CRITERIOS]), 'E': ('Resultado', RESULTADOS), 'F': ('Severidad', SEVERIDAD),
             'G': ('Estado', ESTADOS), 'H': ('Aplica', APLICA), 'I': ('Responsable', ['Ellucian', 'USS', 'Ellucian/USS'])}
    refs = {}
    for col, (name, vals) in lists.items():
        ls[f'{col}1'] = name
        ls[f'{col}1'].font = BOLD
        for i, v in enumerate(vals, 2):
            ls[f'{col}{i}'] = v
            ls[f'{col}{i}'].font = BODY
        refs[name] = f"=Listas!${col}$2:${col}${len(vals) + 1}"

    # ------------------------------------------------------------ Instrucciones
    ins = wb.create_sheet('Instrucciones', 0)
    title(ins, 'Planilla de revisión · Migración R2 · Centros Empresariales',
          'Universidad Señor de Sipán · Elaborado por: Dagner Anibal Chuman Lluen · Septiembre 2026')
    ins.column_dimensions['A'].width = 26
    ins.column_dimensions['B'].width = 95
    rows = [
        ('Objetivo', 'Consolidar los datos e IDs revisados directamente en Banner Student, registrar cada error en el Issue log '
                     'y calcular el porcentaje de correctitud para presentarlo al comité (responsabilidades de la USS, Ellucian diap. 11).'),
        ('1. Revisión', 'Una fila por dato revisado: participante (ID), plantilla, página, campo, valor en el sistema anterior, '
                        'valor en Banner, criterio y resultado. Las celdas amarillas se llenan; las listas desplegables evitan errores.'),
        ('2. Issue log', 'Cada resultado «No conforme» (y cada «Observado» que requiera corrección) se registra aquí con su N° '
                         '(ej. ISS-001) y ese N° se anota en la hoja Revisión. La definición final del Issue log se acordará con Ellucian.'),
        ('3. Resumen', 'Se calcula solo: % de correctitud por criterio, plantilla y centro, comparado con los umbrales de aceptación de Ellucian.'),
        ('4. Cronograma', 'Fechas de carga y revisión de la Migración R2 y horas estimadas para Centros Empresariales según la muestra.'),
        ('5. Tiempos', 'Minutos de validación por participante (Ellucian diap. 9) y si aplica a Centros Empresariales.'),
        ('% correctitud', 'Conformes ÷ revisados. Los «Observado» no suman como correctos hasta que la excepción esté documentada, '
                          'justificada y aprobada por los responsables (criterio de aceptación de Ellucian).'),
        ('Resultados', 'Conforme: igual al sistema anterior y funciona. Observado: diferencia menor o explicada. '
                       'No conforme: falta, se duplica o bloquea un proceso.'),
        ('Leyenda', 'Celdas con fondo amarillo = datos a ingresar. Celdas blancas con fórmula = no modificar.'),
    ]
    r = 4
    for k, v in rows:
        ins.cell(row=r, column=1, value=k).font = BOLD
        c = ins.cell(row=r, column=2, value=v)
        c.font, c.alignment = BODY, WRAP
        ins.cell(row=r, column=1).alignment = WRAP
        ins.row_dimensions[r].height = 32
        r += 1
    ins.cell(row=r - 1, column=2).fill = PatternFill('solid', fgColor=INPUT)
    r += 1
    ins.cell(row=r, column=1, value='Ejemplo de fila (hoja Revisión)').font = Font(name=F, size=11, bold=True, color=PURPLE)
    r += 1
    ex = [('Fecha', '25/09/2026'), ('Revisor', 'D. Chuman'), ('Centro', 'Idiomas'), ('ID Banner', 'A00012345'),
          ('Nombre', 'PÉREZ RAMÍREZ, ANA'), ('Plantilla', '01 Persona'), ('Página', 'SPAIDEN'), ('Campo revisado', 'Segundo apellido'),
          ('Valor sistema anterior', 'RAMÍREZ'), ('Valor en Banner', 'RAMIREZ'), ('Criterio', 'Exactitud'),
          ('Resultado', 'No conforme'), ('N° Issue', 'ISS-001'), ('Observación', 'Se perdió la tilde en la carga.')]
    for h, txt in (('Columna', 'Ejemplo'),):
        for col, v in ((1, h), (2, txt)):
            hc = ins.cell(row=r, column=col, value=v)
            hc.font, hc.fill, hc.alignment, hc.border = HEAD_FONT, HEAD_FILL, CENTER, BORDER
    for i, (k, v) in enumerate(ex, 1):
        kc = ins.cell(row=r + i, column=1, value=k)
        body_cell(kc, font=BOLD)
        vc = ins.cell(row=r + i, column=2, value=v)
        body_cell(vc, fill=GRAY_L)
        vc.font = Font(name=F, size=10, italic=True)
    ins.cell(row=r + len(ex) + 1, column=1, value='Datos ficticios; esta fila de ejemplo no se cuenta en el Resumen.').font = Font(name=F, size=9, italic=True, color='595959')
    ins.sheet_view.showGridLines = False

    # ------------------------------------------------------------ Revisión
    rv = wb.create_sheet('Revisión', 1)
    title(rv, 'Revisión de datos migrados en Banner Student', 'Una fila por dato revisado. Llene las celdas amarillas.')
    heads = ['N°', 'Fecha', 'Revisor', 'Centro', 'ID Banner', 'Nombre', 'Plantilla', 'Página', 'Campo revisado',
             'Valor sistema anterior', 'Valor en Banner', 'Criterio', 'Resultado', 'N° Issue', 'Observación']
    widths = [6, 12, 16, 15, 13, 26, 24, 16, 22, 22, 22, 20, 14, 11, 34]
    style_header(rv, 4, heads, widths)
    for i in range(REV_ROWS):
        row = 5 + i
        n = rv.cell(row=row, column=1, value=f'=IF(E{row}="","",ROW()-4)')
        body_cell(n, align=CENTER)
        for col in range(2, 16):
            c = rv.cell(row=row, column=col)
            body_cell(c, fill=INPUT, fmt='dd/mm/yyyy' if col == 2 else None)
    last = 4 + REV_ROWS
    add_list_validation(rv, f'D5:D{last}', refs['Centro'], 'Centro empresarial')
    add_list_validation(rv, f'G5:G{last}', refs['Plantilla'], 'Plantilla de migración')
    add_list_validation(rv, f'H5:H{last}', refs['Página'], 'Página de Banner donde se revisó')
    add_list_validation(rv, f'L5:L{last}', refs['Criterio'], 'Criterio de aceptación')
    add_list_validation(rv, f'M5:M{last}', refs['Resultado'], 'Resultado de la revisión')
    result_colors(rv, f'M5:M{last}')
    rv.freeze_panes = 'F5'
    rv.auto_filter.ref = f'A4:O{last}'

    # ------------------------------------------------------------ Issue log
    il = wb.create_sheet('Issue log', 2)
    title(il, 'Issue log · errores encontrados', 'Propuesta de campos; la definición final se acordará con Ellucian.')
    heads = ['N° Issue', 'Fecha', 'Centro', 'ID Banner', 'Plantilla', 'Página', 'Descripción del error', 'Criterio',
             'Severidad', 'Evidencia (captura)', 'Reportado por', 'Responsable', 'Estado', 'Fecha de cierre', 'Comentario']
    widths = [11, 12, 15, 13, 24, 16, 40, 20, 11, 24, 16, 14, 14, 13, 30]
    style_header(il, 4, heads, widths)
    for i in range(ISS_ROWS):
        row = 5 + i
        for col in range(1, 16):
            body_cell(il.cell(row=row, column=col), fill=INPUT, fmt='dd/mm/yyyy' if col in (2, 14) else None)
    last_i = 4 + ISS_ROWS
    add_list_validation(il, f'C5:C{last_i}', refs['Centro'], 'Centro empresarial')
    add_list_validation(il, f'E5:E{last_i}', refs['Plantilla'], 'Plantilla de migración')
    add_list_validation(il, f'F5:F{last_i}', refs['Página'], 'Página de Banner')
    add_list_validation(il, f'H5:H{last_i}', refs['Criterio'], 'Criterio afectado')
    add_list_validation(il, f'I5:I{last_i}', refs['Severidad'], 'Alta: bloquea un proceso · Media: dato incorrecto · Baja: formato')
    add_list_validation(il, f'L5:L{last_i}', refs['Responsable'], 'Quién corrige')
    add_list_validation(il, f'M5:M{last_i}', refs['Estado'], 'Estado del issue')
    for val, color in (('Abierto', RED_L), ('En corrección', AMBER_L), ('Resuelto', GREEN_L), ('Cerrado', GRAY_L)):
        il.conditional_formatting.add(f'M5:M{last_i}', CellIsRule(operator='equal', formula=[f'"{val}"'],
                                                                  fill=PatternFill('solid', fgColor=color)))
    il.freeze_panes = 'D5'
    il.auto_filter.ref = f'A4:O{last_i}'

    # ------------------------------------------------------------ Resumen
    rs = wb.create_sheet('Resumen', 3)
    title(rs, 'Resumen para el comité · % de correctitud', 'Se calcula automáticamente con las hojas Revisión e Issue log.')
    R = f"Revisión!$M$5:$M${last}"
    CRIT = f"Revisión!$L$5:$L${last}"
    PLANT = f"Revisión!$G$5:$G${last}"
    CEN = f"Revisión!$D$5:$D${last}"
    for col, w in zip('ABCDEFGH', [30, 12, 12, 12, 12, 14, 14, 14]):
        rs.column_dimensions[col].width = w

    def block(r0, label, items, key_rng, with_threshold):
        heads = [label, 'Revisados', 'Conformes', 'Observados', 'No conformes', '% correctitud']
        heads += ['Umbral', 'Resultado'] if with_threshold else ['Issues abiertos']
        style_header(rs, r0, heads, [None] * len(heads))
        for j, it in enumerate(items):
            rr = r0 + 1 + j
            name = it[0] if isinstance(it, tuple) else it
            rs.cell(row=rr, column=1, value=name)
            rs.cell(row=rr, column=2, value=f'=COUNTIFS({key_rng},$A{rr},{R},"<>")')
            rs.cell(row=rr, column=3, value=f'=COUNTIFS({key_rng},$A{rr},{R},"Conforme")')
            rs.cell(row=rr, column=4, value=f'=COUNTIFS({key_rng},$A{rr},{R},"Observado")')
            rs.cell(row=rr, column=5, value=f'=COUNTIFS({key_rng},$A{rr},{R},"No conforme")')
            rs.cell(row=rr, column=6, value=f'=IF(B{rr}=0,"",C{rr}/B{rr})')
            if with_threshold:
                rs.cell(row=rr, column=7, value=it[2])
                rs.cell(row=rr, column=8, value=f'=IF(B{rr}=0,"Sin datos",IF(F{rr}>=G{rr},"Cumple","No cumple"))')
                rs.cell(row=rr, column=7).comment = Comment('Umbral sugerido por Ellucian (Criterios de aceptación, diap. 10).', 'USS')
            else:
                col = 'E' if label == 'Plantilla' else 'C'
                rs.cell(row=rr, column=7, value=f"=COUNTIFS('Issue log'!${col}$5:${col}${last_i},$A{rr},"
                                                f"'Issue log'!$M$5:$M${last_i},\"Abierto\")+COUNTIFS('Issue log'!${col}$5:${col}${last_i},$A{rr},"
                                                f"'Issue log'!$M$5:$M${last_i},\"En corrección\")")
            for cc in range(1, len(heads) + 1):
                body_cell(rs.cell(row=rr, column=cc), align=CENTER if cc > 1 else WRAP, font=BOLD if cc == 1 else BODY)
            rs.cell(row=rr, column=6).number_format = '0.0%'
            if with_threshold:
                rs.cell(row=rr, column=7).number_format = '0.0%'
                rs.cell(row=rr, column=7).font = Font(name=F, size=10, color='0000FF')
        end = r0 + len(items)
        tr = end + 1
        rs.cell(row=tr, column=1, value='Total')
        for cc, colL in ((2, 'B'), (3, 'C'), (4, 'D'), (5, 'E')):
            rs.cell(row=tr, column=cc, value=f'=SUM({colL}{r0 + 1}:{colL}{end})')
        rs.cell(row=tr, column=6, value=f'=IF(B{tr}=0,"",C{tr}/B{tr})')
        if not with_threshold:
            rs.cell(row=tr, column=7, value=f'=SUM(G{r0 + 1}:G{end})')
        for cc in range(1, len(heads) + 1):
            body_cell(rs.cell(row=tr, column=cc), fill=PURPLE_L, align=CENTER if cc > 1 else WRAP, font=BOLD)
        rs.cell(row=tr, column=6).number_format = '0.0%'
        if with_threshold:
            rng = f'H{r0 + 1}:H{end}'
            rs.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"Cumple"'], fill=PatternFill('solid', fgColor=GREEN_L)))
            rs.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"No cumple"'], fill=PatternFill('solid', fgColor=RED_L)))
        return tr + 2

    r = block(4, 'Criterio', CRITERIOS, CRIT, True)
    rs.cell(row=r - 1, column=1, value='% correctitud = Conformes ÷ Revisados. Umbrales en azul: sugeridos por Ellucian; ajústelos solo si el comité lo aprueba.').font = Font(name=F, size=9, italic=True, color='595959')
    r = block(r + 1, 'Plantilla', PLANTILLAS, PLANT, False)
    r = block(r, 'Centro', CENTROS, CEN, False)
    # issues
    style_header(rs, r, ['Issue log', 'Total', 'Abiertos', 'En corrección', 'Resueltos', 'Cerrados', 'Severidad alta'], [None] * 7)
    IL = f"'Issue log'!$M$5:$M${last_i}"
    vals = [f"=COUNTA('Issue log'!$A$5:$A${last_i})", f'=COUNTIF({IL},"Abierto")', f'=COUNTIF({IL},"En corrección")',
            f'=COUNTIF({IL},"Resuelto")', f'=COUNTIF({IL},"Cerrado")', f"=COUNTIF('Issue log'!$I$5:$I${last_i},\"Alta\")"]
    rs.cell(row=r + 1, column=1, value='Issues registrados')
    for j, v in enumerate(vals, 2):
        rs.cell(row=r + 1, column=j, value=v)
    for cc in range(1, 8):
        body_cell(rs.cell(row=r + 1, column=cc), align=CENTER if cc > 1 else WRAP, font=BOLD if cc == 1 else BODY)
    rs.sheet_view.showGridLines = False

    # ------------------------------------------------------------ Tiempos
    tp = wb.create_sheet('Tiempos', 4)
    title(tp, 'Tiempos de validación por participante', 'Fuente: Ellucian, «Tiempos de validación por alumnos» (diap. 9). Columna «Aplica a CE» editable.')
    heads = ['Plantilla', '¿Qué se espera validar?', 'Página', 'Min (Ellucian)', 'Etapa del cronograma', 'Aplica a CE', 'Min CE']
    style_header(tp, 4, heads, [24, 60, 20, 12, 24, 14, 10])
    for j, (pl, q, pg, m, et, ap) in enumerate(TIEMPOS):
        rr = 5 + j
        vals = [pl, q, pg, m, et, ap, f'=IF(F{rr}="No",0,D{rr})']
        for cc, v in enumerate(vals, 1):
            c = tp.cell(row=rr, column=cc, value=v)
            body_cell(c, fill=INPUT if cc == 6 else None, align=CENTER if cc in (4, 6, 7) else WRAP)
        tp.cell(row=rr, column=4).font = Font(name=F, size=10, color='0000FF')
    tl = 4 + len(TIEMPOS)
    add_list_validation(tp, f'F5:F{tl}', refs['Aplica'], '¿Aplica a Centros Empresariales?')
    for rr, lab, formula in ((tl + 1, 'Minutos totales', f'=SUM(D5:D{tl})'), (tl + 2, 'Horas', f'=D{tl + 1}/60')):
        tp.cell(row=rr, column=3, value=lab)
        tp.cell(row=rr, column=4, value=formula)
        tp.cell(row=rr, column=7, value=f'=SUM(G5:G{tl})' if rr == tl + 1 else f'=G{tl + 1}/60')
        for cc in (3, 4, 7):
            body_cell(tp.cell(row=rr, column=cc), fill=PURPLE_L, align=CENTER, font=BOLD)
        tp.cell(row=rr, column=4).number_format = tp.cell(row=rr, column=7).number_format = '0.0' if rr == tl + 2 else '0'
    tp.cell(row=tl + 3, column=1, value='«Por confirmar» se cuenta en Min CE (estimación conservadora). El tiempo de SIAINST es por docente, no por participante.').font = Font(name=F, size=9, italic=True, color='595959')
    tp.freeze_panes = 'A5'

    # ------------------------------------------------------------ Cronograma
    cr = wb.create_sheet('Cronograma', 4)
    title(cr, 'Cronograma Migración R2 · carga y revisión', 'Fuente: Ellucian, «Migración R2 - Tiempos de validación - Cronograma» (diap. 12).')
    cr['A3'], cr['C3'] = 'Participantes por centro (muestra)', 16
    cr['D3'], cr['F3'] = 'Centros', 3
    cr['G3'], cr['I3'] = 'Docentes (muestra)', 10
    cr['J3'], cr['L3'] = 'Egresados por centro', 5
    for ref in ('C3', 'F3', 'I3', 'L3'):
        cr[ref].fill = PatternFill('solid', fgColor=INPUT)
        cr[ref].font = Font(name=F, size=10, bold=True, color='0000FF')
        cr[ref].border = BORDER
        cr[ref].alignment = CENTER
    for ref in ('A3', 'D3', 'G3', 'J3'):
        cr[ref].font = Font(name=F, size=9, bold=True)
    cr['C3'].comment = Comment('Propuesta: 16 casos por programa (ejemplo de Ellucian).', 'USS')
    heads = ['Etapa', 'Ambiente', 'Carga inicio', 'Carga fin', 'Revisión inicio', 'Revisión fin', 'Aplica a CE',
             'Min por caso', 'Muestra CE', 'Horas CE', 'Plazo (según hoy)', 'Estado de la revisión CE']
    style_header(cr, 5, heads, [30, 10, 12, 12, 14, 13, 13, 11, 11, 10, 16, 20])
    for j, (et, amb, ci, cf, ri, rf, mt) in enumerate(CRONO):
        rr = 6 + j
        aplica = f'=IFERROR(INDEX(Tiempos!$F$5:$F${tl},MATCH(A{rr},Tiempos!$E$5:$E${tl},0)),"—")'
        minutos = f'=IF(G{rr}="—","",SUMIFS(Tiempos!$G$5:$G${tl},Tiempos!$E$5:$E${tl},A{rr}))'
        muestra = {'part': '=$C$3*$F$3', 'doc': '=$I$3', 'egr': '=$L$3*$F$3', None: ''}[mt]
        horas = f'=IF(OR(H{rr}="",I{rr}=""),"",H{rr}*I{rr}/60)'
        plazo = (f'=IF(E{rr}="","",IF(G{rr}="No","No aplica",IF(TODAY()<E{rr},"Próxima",IF(TODAY()<=F{rr},"En revisión","Plazo cerrado"))))')
        vals = [et, amb, ci, cf, ri, rf, aplica, minutos, muestra, horas, plazo, None]
        for cc, v in enumerate(vals, 1):
            c = cr.cell(row=rr, column=cc, value=v)
            body_cell(c, fill=INPUT if cc == 12 else None, align=CENTER if cc > 1 else WRAP,
                      fmt='dd/mm/yyyy' if cc in (3, 4, 5, 6) else ('0.0' if cc == 10 else None))
        cr.cell(row=rr, column=1).font = BOLD
    cl = 5 + len(CRONO)
    cr.cell(row=cl + 1, column=9, value='Total horas')
    cr.cell(row=cl + 1, column=10, value=f'=SUM(J6:J{cl})')
    for cc in (9, 10):
        body_cell(cr.cell(row=cl + 1, column=cc), fill=PURPLE_L, align=CENTER, font=BOLD, fmt='0.0' if cc == 10 else None)
    add_list_validation(cr, f'L6:L{cl}', '"Pendiente,En curso,Terminada,No aplica"', 'Estado de la revisión de Centros Empresariales')
    for val, color in (('En revisión', AMBER_L), ('Próxima', PURPLE_L), ('Plazo cerrado', GRAY_L)):
        cr.conditional_formatting.add(f'K6:K{cl}', CellIsRule(operator='equal', formula=[f'"{val}"'], fill=PatternFill('solid', fgColor=color)))
    cr.cell(row=cl + 3, column=1, value='Horas CE = minutos por caso (hoja Tiempos) × muestra ÷ 60. «Plazo» se actualiza con la fecha del día. '
                                        'PROD: Personas a Matriz de seguridad; TEST: desde Estudiantes (después de la clonación).').font = Font(name=F, size=9, italic=True, color='595959')
    cr.freeze_panes = 'B6'

    # orden y cierre
    wb.move_sheet('Listas', offset=len(wb.sheetnames))
    ls.sheet_state = 'hidden'
    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = PURPLE if ws.title in ('Revisión', 'Issue log') else '4EA72E'
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.fitToWidth = 1
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ws.page_setup.fitToHeight = 0
    wb.active = 0
    wb.properties.creator = 'Dagner Anibal Chuman Lluen'
    wb.properties.title = 'Planilla de revisión - Migración R2 - Centros Empresariales'
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    print('OK', OUT)


if __name__ == '__main__':
    build()
