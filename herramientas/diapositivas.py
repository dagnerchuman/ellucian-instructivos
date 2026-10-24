#!/usr/bin/env python3
"""Convierte los instructivos PPTX en imágenes WebP, una por diapositiva.

El visor (visor_instructivos/) y las guías paso a paso de los diagramas muestran estas
imágenes en lugar de leer el PPTX en el navegador.

    python3 herramientas/diapositivas.py                 # convierte los que faltan o cambiaron
    python3 herramientas/diapositivas.py --solo 5.3_4.1.4.1.6 --solo 6.2.1_Asignar
    python3 herramientas/diapositivas.py --forzar        # vuelve a convertir todo

Salida:
    visor_instructivos/diapositivas/<id>/NNN.webp           diapositiva completa
    visor_instructivos/diapositivas/<id>/recortes/NNN.webp  recorte sin título ni banda lateral (guías)
    visor_instructivos/diapositivas/indice.js               window.DIAPOSITIVAS = {id: {n, ancho, alto, bytes}}

Necesita LibreOffice (soffice), pdftoppm (poppler) y Pillow con WebP. Los <id> son los de
visor_instructivos/data.js; las carpetas CAPACIDAD n/ no se mueven (ADR-015).
"""
import argparse
import json
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from PIL import Image

RAIZ = Path(__file__).resolve().parents[1]
DATA_JS = RAIZ / "visor_instructivos" / "data.js"
SALIDA = RAIZ / "visor_instructivos" / "diapositivas"
INDICE = SALIDA / "indice.js"
ANCHO = 1600
CALIDAD = 72
# Proporciones que se quitan en los recortes: banda lateral izquierda y título superior (diseño Ellucian).
CAJA_RECORTE = (0.034, 0.118, 0.0, 0.075)
VERSION_RECORTE = 2  # súbelo si cambia la forma de recortar: obliga a rehacer los recortes


def presentaciones():
    """{id: {path, sizeBytes, title…}} leído de visor_instructivos/data.js."""
    texto = DATA_JS.read_text(encoding="utf-8")
    marca = "var PRESENTATIONS = window.PRESENTATIONS = "
    i = texto.index(marca) + len(marca)
    j = texto.index("\n};", i) + 2
    return json.loads(texto[i:j])


def leer_indice():
    if not INDICE.exists():
        return {}
    texto = INDICE.read_text(encoding="utf-8")
    return json.loads(texto[texto.index("{"):texto.rindex("}") + 1])


def escribir_indice(indice):
    ordenado = dict(sorted(indice.items()))
    INDICE.write_text(
        "// Archivo generado por herramientas/diapositivas.py: no lo edites a mano.\n"
        "window.DIAPOSITIVAS = " + json.dumps(ordenado, ensure_ascii=False, indent=1) + ";\n",
        encoding="utf-8",
    )


def nombre_diapositiva(n):
    return f"{n:03d}.webp"


def convertir_uno(id_, pres, tmp):
    pptx = RAIZ / pres["path"]
    destino = SALIDA / id_
    perfil = Path(tmp) / "perfil"
    pdf_dir = Path(tmp) / "pdf"
    png_dir = Path(tmp) / "png"
    for d in (pdf_dir, png_dir):
        shutil.rmtree(d, ignore_errors=True)
        d.mkdir(parents=True)
    subprocess.run(
        ["soffice", f"-env:UserInstallation=file://{perfil}", "--headless", "--norestore",
         "--convert-to", "pdf", "--outdir", str(pdf_dir), str(pptx)],
        check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=600,
    )
    pdf = next(pdf_dir.glob("*.pdf"))
    subprocess.run(["pdftoppm", "-scale-to-x", str(ANCHO), "-scale-to-y", "-1", "-png", str(pdf), str(png_dir / "d")],
                   check=True, timeout=600)
    pngs = sorted(png_dir.glob("d-*.png"), key=lambda p: int(p.stem.split("-")[-1]))
    shutil.rmtree(destino, ignore_errors=True)
    destino.mkdir(parents=True)
    ancho = alto = 0
    for n, png in enumerate(pngs, start=1):
        with Image.open(png) as im:
            ancho, alto = im.size
            im.convert("RGB").save(destino / nombre_diapositiva(n), "WEBP", quality=CALIDAD, method=6)
    return {"n": len(pngs), "ancho": ancho, "alto": alto, "bytes": pptx.stat().st_size}


def convertir(filtros=None, forzar=False):
    if not shutil.which("soffice") or not shutil.which("pdftoppm"):
        sys.exit("Faltan soffice (LibreOffice) o pdftoppm (poppler).")
    SALIDA.mkdir(parents=True, exist_ok=True)
    pres = presentaciones()
    indice = leer_indice()
    elegidos = [i for i in pres if not filtros or any(f in i for f in filtros)]
    hechos = 0
    with tempfile.TemporaryDirectory() as tmp:
        for k, id_ in enumerate(elegidos, start=1):
            previo = indice.get(id_)
            completo = previo and previo.get("bytes") == (RAIZ / pres[id_]["path"]).stat().st_size and \
                len(list((SALIDA / id_).glob("*.webp"))) == previo.get("n")
            if completo and not forzar:
                continue
            try:
                indice[id_] = convertir_uno(id_, pres[id_], tmp)
            except (subprocess.SubprocessError, StopIteration, OSError) as e:
                print(f"FALLA  {id_}: {e}")
                continue
            hechos += 1
            print(f"OK     [{k}/{len(elegidos)}] {id_}: {indice[id_]['n']} diapositivas")
            escribir_indice(indice)
    escribir_indice(indice)
    print(f"\n{hechos} instructivos convertidos; {len(indice)} en el índice.")


_CAJAS = {}


def _cajas_presentacion(id_):
    """Por diapositiva: ((izquierda, arriba, derecha, abajo) del contenido, títulos quitados).

    La caja deja fuera el título, la banda lateral y el pie. Los títulos quitados se guardan como
    (x0, y0, x1, y1) en fracciones de la lámina: recorte() los vuelve a meter si el borde los parte."""
    if id_ in _CAJAS:
        return _CAJAS[id_]
    from pptx import Presentation
    from pptx.enum.shapes import PP_PLACEHOLDER
    pres = Presentation(RAIZ / presentaciones()[id_]["path"])
    W, H = pres.slide_width, pres.slide_height
    textos = _textos_presentacion(id_)
    cajas = []
    for k, lamina in enumerate(pres.slides):
        conservados = {p["texto"] for p in textos[k]}
        titulos = []
        x0 = y0 = None
        x1 = y1 = 0
        for forma in lamina.shapes:
            if forma.left is None or forma.top is None or not forma.width or not forma.height:
                continue
            izq, sup, ancho, alto = forma.left, forma.top, forma.width, forma.height
            if ancho * alto > 0.85 * W * H:  # fondo
                continue
            if sup > 0.9 * H or (sup + alto > 0.96 * H and alto < 0.08 * H):  # pie: logo, derechos, número
                continue
            if izq + ancho < 0.06 * W and alto > 0.4 * H:  # banda lateral
                continue
            caja_forma = (izq / W, sup / H, (izq + ancho) / W, (sup + alto) / H)
            if forma.is_placeholder and forma.placeholder_format.type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
                titulos.append(caja_forma)
                continue
            if forma.has_text_frame:
                texto = " ".join(forma.text_frame.text.split())
                if not texto:
                    pass
                elif sup < 0.16 * H and sup + alto < 0.26 * H and len(texto) < 120:
                    titulos.append(caja_forma)
                    continue  # título de la lámina
                elif not any(c[:25] in texto for c in conservados) and len(texto) > 3 and not texto.isdigit():
                    continue  # texto repetido del diseño
            x0 = izq if x0 is None else min(x0, izq)
            y0 = sup if y0 is None else min(y0, sup)
            x1 = max(x1, izq + ancho)
            y1 = max(y1, sup + alto)
        if x0 is None:
            cajas.append((None, titulos))
            continue
        margen_x, margen_y = 0.008 * W, 0.012 * H
        cajas.append(((max(0, (x0 - margen_x) / W), max(0, (y0 - margen_y) / H),
                       max(0, 1 - (x1 + margen_x) / W), max(0, 1 - (y1 + margen_y) / H)), titulos))
    _CAJAS[id_] = cajas
    return cajas


def caja_contenido(id_, n):
    """Caja de recorte de la diapositiva n según la posición de sus elementos; CAJA_RECORTE si no se puede."""
    try:
        cajas = _cajas_presentacion(id_)
        caja = cajas[n - 1][0] if 0 < n <= len(cajas) else None
    except Exception:  # noqa: BLE001 - si el PPTX no se puede leer, se usa el recorte fijo
        caja = None
    if not caja or (1 - caja[0] - caja[2]) < 0.3 or (1 - caja[1] - caja[3]) < 0.3:
        return CAJA_RECORTE
    return caja


def titulos_quitados(id_, n):
    """Cajas (x0, y0, x1, y1), en fracciones, de los títulos que caja_contenido() dejó fuera."""
    try:
        cajas = _cajas_presentacion(id_)
        return cajas[n - 1][1] if 0 < n <= len(cajas) else []
    except Exception:  # noqa: BLE001
        return []


def _con_tinta(gris, caja, minimo=3):
    """True si en esa franja de la imagen hay trazos (letras), no solo fondo."""
    x0, y0, x1, y1 = caja
    if x1 <= x0 or y1 <= y0:
        return False
    return sum(gris.crop(caja).histogram()[:200]) >= minimo


def _sin_titulos_partidos(gris, x0, y0, x1, y1, titulos):
    """Evita que un borde del recorte parta las letras de un título.

    Si el borde de arriba solo roza el pie de las letras, baja hasta la primera fila limpia (sin pasar
    del margen, donde empieza el contenido). Si no hay fila limpia o el corte es por un lado, el
    recorte se agranda para que el título entre entero."""
    w, h = gris.size
    mx, my = round(0.008 * w), round(0.012 * h)
    for tx0, ty0, tx1, ty1 in titulos:
        tx0, ty0, tx1, ty1 = round(w * tx0), round(h * ty0), round(w * tx1), round(h * ty1)
        if tx1 <= x0 or tx0 >= x1 or ty1 <= y0 or ty0 >= y1:
            continue  # el título queda fuera del recorte
        cols, filas = (max(x0, tx0), min(x1, tx1)), (max(y0, ty0), min(y1, ty1))
        if ty0 < y0 < ty1 and _con_tinta(gris, (cols[0], y0 - 1, cols[1], y0 + 2)):
            limpia = next((y for y in range(y0, y0 + my) if not _con_tinta(gris, (cols[0], y, cols[1], y + 1), 1)), None)
            if limpia is not None:
                y0 = limpia
                continue
            partido = True
        else:
            partido = ((tx0 < x0 < tx1 and _con_tinta(gris, (x0 - 1, filas[0], x0 + 2, filas[1])))
                       or (tx0 < x1 < tx1 and _con_tinta(gris, (x1 - 2, filas[0], x1 + 1, filas[1]))))
        if partido:
            x0, y0 = max(0, min(x0, tx0 - mx)), max(0, min(y0, ty0 - my))
            x1 = min(w, max(x1, tx1 + mx))
    return x0, y0, x1, y1


def _extender_abajo(im, x0, x1, y1, limite):
    """Baja el borde inferior mientras haya contenido (un recuadro con texto que desborda su forma)."""
    gris = im.convert("L")
    vacias = 0
    y = y1
    while y < limite and vacias < 14:
        fila = gris.crop((x0, y, x1, y + 1)).histogram()
        oscuros = sum(fila[:235])
        if oscuros > 0.02 * (x1 - x0):
            vacias = 0
            y1 = y + 1
        else:
            vacias += 1
        y += 1
    return min(y1 + 6, limite)


def recorte(id_, n, caja=None):
    """Ruta (relativa a la raíz) del recorte de la diapositiva n; lo crea si falta. None si no existe la diapositiva."""
    origen = SALIDA / id_ / nombre_diapositiva(n)
    if not origen.exists():
        return None
    destino = SALIDA / id_ / "recortes" / nombre_diapositiva(n)
    if caja:  # caja propia de un paso: archivo aparte, para no pisar el recorte automático que usan otros pasos
        destino = destino.with_name(f"{destino.stem}-{'-'.join(f'{round(v * 1000):03d}' for v in caja)}{destino.suffix}")
    titulos = [] if caja else titulos_quitados(id_, n)
    caja = tuple(caja or caja_contenido(id_, n))
    marca = destino.with_suffix(".caja")
    firma = f"v{VERSION_RECORTE}:" + ",".join(f"{v:.4f}" for v in caja)
    vigente = destino.exists() and marca.exists() and marca.read_text() == firma
    if not vigente or destino.stat().st_mtime < origen.stat().st_mtime:
        destino.parent.mkdir(exist_ok=True)
        with Image.open(origen) as im:
            w, h = im.size
            izq, sup, der, inf = caja
            x0, y0, x1, y1 = round(w * izq), round(h * sup), round(w * (1 - der)), round(h * (1 - inf))
            x0, y0, x1, y1 = _sin_titulos_partidos(im.convert("L"), x0, y0, x1, y1, titulos)
            y1 = _extender_abajo(im, x0, x1, y1, round(h * 0.95))
            im.crop((x0, y0, x1, y1)).save(destino, "WEBP", quality=CALIDAD, method=6)
        marca.write_text(firma)
    return destino.relative_to(RAIZ).as_posix()


_TEXTOS = {}


def _textos_presentacion(id_):
    """Por diapositiva: lista de párrafos [{"num", "texto", "nivel"}] sin el texto que se repite en el diseño."""
    if id_ in _TEXTOS:
        return _TEXTOS[id_]
    from collections import Counter
    from pptx import Presentation
    from pptx.oxml.ns import qn
    pres = Presentation(RAIZ / presentaciones()[id_]["path"])
    por_lamina = []
    for lamina in pres.slides:
        parrafos = []
        for forma in lamina.shapes:
            if not forma.has_text_frame:
                continue
            contador = None
            for p in forma.text_frame.paragraphs:
                texto = " ".join("".join(r.text for r in p.runs).split())
                ppr = p._p.pPr
                auto = ppr.find(qn("a:buAutoNum")) if ppr is not None else None
                if auto is not None:
                    contador = (contador + 1) if contador is not None else int(auto.get("startAt", "1"))
                num = str(contador) if auto is not None else ""
                escrito = re.match(r"^(\d{1,2})[.)]\s*(\S.*)$", texto)
                if not num and escrito:  # número escrito a mano en la lámina («20.Guardar la información»)
                    num, texto = escrito.group(1), escrito.group(2)
                if texto:
                    parrafos.append({"num": num, "texto": texto, "nivel": p.level})
        por_lamina.append(parrafos)
    # Lo que aparece en muchas láminas es diseño (banda lateral, pie, logo): se descarta.
    frecuencia = Counter(p["texto"] for ps in por_lamina for p in {x["texto"]: x for x in ps}.values())
    limite = max(3, len(por_lamina) * 0.3)
    es_titulo = re.compile(r"^[^.]{3,50}\([A-Z]{6,8}\)$")  # «Programar NRC (SSASECT)»: título de la lámina
    _TEXTOS[id_] = [[p for p in ps if frecuencia[p["texto"]] < limite and len(p["texto"]) > 3
                     and not p["texto"].isdigit() and not (not p["num"] and es_titulo.match(p["texto"]))]
                    for ps in por_lamina]
    return _TEXTOS[id_]


def texto_diapositiva(id_, n):
    """Párrafos de la diapositiva n (1, 2…) para mostrarlos junto a la imagen."""
    try:
        laminas = _textos_presentacion(id_)
    except Exception:  # noqa: BLE001 - el texto es un complemento; sin él se muestra solo la imagen
        return []
    return laminas[n - 1] if 0 < n <= len(laminas) else []


def diapositiva(id_, n):
    """Ruta (relativa a la raíz) de la diapositiva completa, o None si no se ha convertido."""
    ruta = SALIDA / id_ / nombre_diapositiva(n)
    return ruta.relative_to(RAIZ).as_posix() if ruta.exists() else None


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--solo", action="append", help="convierte solo los id que contengan este texto (se puede repetir)")
    ap.add_argument("--forzar", action="store_true", help="vuelve a convertir aunque ya existan")
    args = ap.parse_args()
    convertir(args.solo, args.forzar)
