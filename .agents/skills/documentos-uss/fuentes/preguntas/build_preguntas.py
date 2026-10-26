"""Word de preguntas para la reunión con Ellucian (y las dudas de la USS), con un recuadro de respuesta por pregunta.

Los textos salen del registro de dudas (`registro.md`), por su ID; aquí solo se agrupan, se titulan y se agrega
«Lo que ya sabemos». Para leer las respuestas escritas en el Word: `leer_respuestas.py`.

    python3 build_preguntas.py "docs/dudas/PREGUNTAS PARA ELLUCIAN - REUNIÓN 27-10 - CENTROS EMPRESARIALES.docx"
"""
import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AQUI = Path(__file__).resolve().parent
SKILL = AQUI.parents[1]
RAIZ = SKILL.parents[2]
REGISTRO = SKILL.parent / "preguntas-y-dudas-centros" / "registro.md"
sys.path.insert(0, str(SKILL / "scripts" / "pdf"))
import common  # noqa: E402  (PAGINAS y SIGLAS: significado de cada código)

AUTOR = "Dagner Anibal Chuman Lluen"
FECHA_REUNION = "27/10/2026"
MORADO, MORADO_OSC, VERDE = "7030A0", "5C2193", "4EA72E"
GRIS = "6B6478"

# Significados que no están en common.py (tomados del registro y de los instructivos).
EXTRA = {
    "GORRSQL": "reglas SQL de los procesos",
    "GORICCR": "valores de configuración de las integraciones",
    "SIAFLCT": "reglas de carga por tipo de contrato",
    "SIACONA": "análisis de la carga por contrato",
    "STVCNTR": "catálogo de reglas de carga del contrato",
    "STVNIST": "catálogo de labores no educativas",
    "STVFCNT": "catálogo de tipos de contrato docente",
    "SAAADMS": "solicitud de admisión",
    "SFRFASC": "estimación de cuotas por lotes",
    "SSADETL": "detalle del NRC: ligas y cuotas",
    "SVVSRVC": "catálogo de solicitudes de servicio",
}

# ---------------------------------------------------------------------------------------------
# Contenido: bloques en orden de impacto (primero lo que bloquea), con título corto y contexto.
# "texto" reemplaza al del registro solo cuando el original es demasiado largo para la reunión.
# ---------------------------------------------------------------------------------------------
ELLUCIAN = [
    ("Notas y asistencia (hoy bloquean el cierre)", [
        dict(id="U32", titulo="La asistencia no llega sola a ATTRGRD",
             texto="El porcentaje de asistencia no llega a **ATTRGRD** aunque la configuración de TI del instructivo 7.1.0 "
                   "está completa (GORICCR en Y, GORRSQL activa y validada). ¿Qué falta revisar: (a) el segundo registro de la "
                   "regla en GORRSQL; (b) el cálculo por horas, si el panel de Pasar lista muestra «Horas asistidas 00:00» con la "
                   "unidad Presente/Ausente de SOAATRK; (c) que el tipo de reunión y la sesión coincidan con `:MEETING_CODE` y "
                   "`:SESSION_ID`; (d) la acción del docente que dispara el envío?",
             sabemos="Con el porcentaje cargado, la nota final sale INH en SFASLST (C105, C107). Según el instructivo 7.1.1 "
                     "(diap. 56 y 57), el docente pulsa Clase › «Actualizar la asistencia en el libro de calificaciones» (R21); "
                     "falta probarlo con el NRC 10044."),
        dict(id="E22", titulo="Escala de Idiomas y Emprendimiento en SHAGRDE",
             sabemos="Idiomas califica con el modo C (00 a 100, aprueba con 70); Computación y Emprendimiento con el modo V "
                     "(00 a 20, aprueba con 11) (C65, ADR-009)."),
    ]),
    ("Periodos, horas y prerrequisitos", [
        dict(id="E01", titulo="Hora de clase de 45 y 50 minutos en el mismo periodo",
             sabemos="SIATERM tiene un solo factor de duración por periodo y SSASECT calcula las horas con ese valor "
                     "(R07; 5.2, diap. 12). Hoy en SEUSS: 45 min de día y 50 de noche cuentan como 1 hora (C05)."),
        dict(id="E03", titulo="Hora de 60 minutos",
             sabemos="Además de la hora de clase existe una hora presencial de 60 min, como la hora de trabajo (C05)."),
        dict(id="E13", titulo="«En progreso» con grupos seguidos",
             sabemos="Con «En progreso» marcado en SOATERM, el curso en curso cuenta como completo para el prerrequisito; "
                     "se revisa contra SSAPREQ al calificar (R10; 1.1.2, diap. 18 y 19)."),
        dict(id="E04", titulo="BASIC I después de pasar con examen de suficiencia",
             sabemos="El prerrequisito «BASIC I aprobado o suficiencia con puntaje mínimo» va en SCAPREQ y el puntaje en "
                     "SOATEST (R03, R09). En Idiomas, quien desaprueba no avanza salvo con suficiencia (C11)."),
    ]),
    ("Mallas y CAPP", [
        dict(id="E05", titulo="Malla de los centros en CAPP",
             sabemos="Reglas curriculares cargadas: Idiomas MI03-01 y MI03-02; Computación MC38-01 (regla ECOM-01); "
                     "Emprendimiento MM01-01 a MM01-03. La matrícula por Proyecciones del Autoservicio ya se validó (C71)."),
        dict(id="E15", titulo="Programa sin plan completo al abrir 20271",
             sabemos="Sin cursos en el catálogo no hay NRC; los cursos nuevos se crean antes de cambiar las áreas y los "
                     "equivalentes deben existir antes (R12)."),
        dict(id="E16", titulo="Lo que ve el tutor sin malla cargada"),
    ]),
    ("Docentes y carga lectiva", [
        dict(id="E14", titulo="Reemplazo de un docente a mitad del curso",
             sabemos="El nuevo docente se asigna en SSASECT como principal con su % de responsabilidad; si PRIMINSTR = Y, "
                     "solo el principal registra notas (R11)."),
        dict(id="E20", titulo="SIAFLRT o SIAFLCT y el factor FTE",
             sabemos="El instructivo las presenta como dos formas de evaluar la carga, no excluyentes (R19). "
                     "ADR-020 recomienda medir por contrato (SIAFLCT y SIACONA). Contratos ES (48 h) y FC ya creados (C54 a C57)."),
    ]),
    ("Finanzas y matrícula", [
        dict(id="U29", titulo="Cambio de modalidad de pago entre partes de periodo",
             sabemos="Pregunta de Sandra Valdiviezo (USS) en el chat de coordinación. El plan de cuotas I1 se probó en TEST "
                     "(TSAISTC, TSPISTA, C76 a C80)."),
        dict(id="E23", titulo="Estimación en línea para pagar al matricular",
             sabemos="Liliana Ocampo (Ellucian) indicó que hay que habilitarla para que el cargo se liquide al instante, "
                     "sin esperar el proceso por lotes SFRFASC (C82)."),
        dict(id="E24", titulo="GAP de Finanzas para elegir el plan de pagos",
             sabemos="La selección del plan por el alumno entrará con un desarrollo a medida pedido por Finanzas (C82). "
                     "En TEST el plan se asignó en SAAADMS (C78)."),
    ]),
    ("Aula virtual", [
        dict(id="E02", titulo="Qué sale al aula virtual y si las notas regresan",
             sabemos="El NRC se vincula con el campo «Socio de integración» de SSASECT; los datos del NRC salen al aula "
                     "virtual (R06; 6.2.4, diap. 10 a 14)."),
    ]),
    ("Solicitudes, certificados y seguimiento", [
        dict(id="E21", titulo="Certificados y documentos del alumno",
             sabemos="El certificado es una solicitud de servicio (CER) que el alumno crea en el Autoservicio y el área "
                     "atiende en SVASVPR; se validó con Computación (C44, R20)."),
        dict(id="E10", titulo="Certificado de los centros como grado"),
        dict(id="E17", titulo="Aviso por correo y derivación a Gobierno de Personas",
             sabemos="Una queja se registra como solicitud de servicio (R14). En SVVSRVC hay 17 servicios y ninguno es de "
                     "queja (C44)."),
        dict(id="E18", titulo="Quién atiende cada servicio y confidencialidad"),
        dict(id="E19", titulo="Alertas tempranas por faltas o notas bajas",
             sabemos="Los instructivos traen seguimiento en SPACMNT, estado académico y comunicaciones, pero no alertas "
                     "tempranas (R16)."),
    ]),
    ("Migración y validación", [
        dict(id="E06", titulo="Historia de los centros en la migración"),
        dict(id="E07", titulo="Carga LD01 con equivalencias",
             sabemos="Ellucian indicó que cada fila de equivalencias debe ir con el plan de su modalidad (C58)."),
        dict(id="E08", titulo="Alumnos de los centros en la muestra de validación"),
        dict(id="E09", titulo="Quién ejecuta SMARQCM y SFPPROJ en TEST"),
        dict(id="E11", titulo="Formato del Issue log"),
        dict(id="E12", titulo="Totales por centro desde Insight"),
    ]),
]

USS = [
    ("Reglas académicas", [
        dict(id="U01", titulo="Puntaje mínimo del examen de suficiencia de Inglés"),
        dict(id="U21", titulo="¿Emprendimiento rinde examen de suficiencia?"),
        dict(id="U02", titulo="Secuencia de cursos en Emprendimiento",
             texto="En Idiomas y Computación ya se definió la secuencia (C62). ¿En **Emprendimiento** el alumno debe "
                   "aprobar el módulo I para llevar el II y el III, o puede llevarlos en cualquier orden?"),
        dict(id="U09", titulo="Nombres de las horas (pedagógica y cronológica)"),
    ]),
    ("Programas, mallas y catálogo", [
        dict(id="U31", titulo="¿CMEMC44 o CMEMC38 para Diseño Web?"),
        dict(id="U19", titulo="Programa MC38 y regla ECOM-01",
             texto="`CMEMC38` y el nivel `C` ya están confirmados. ¿Qué programa es **MC38** en SMAAREA y la regla "
                   "**ECOM-01** es el requisito de computación de pregrado?"),
        dict(id="U05", titulo="Nombres de los cursos de Emprendimiento en el catálogo",
             texto="En Idiomas y Computación los cursos ya están confirmados. ¿Cuáles son los nombres y códigos de los "
                   "cursos de **Emprendimiento** (`ESGE`) en el catálogo de Banner?"),
        dict(id="U12", titulo="Planes en rediseño y de posgrado antes de 20271"),
        dict(id="U13", titulo="Código del periodo «20271»"),
    ]),
    ("Matrícula, permisos y egresados", [
        dict(id="U10", titulo="Quién da sobrepasos de prerrequisito"),
        dict(id="U08", titulo="¿La jefatura autoriza los sobrepasos?",
             texto="El jefe da el visto bueno de la carga lectiva (C47). ¿También autoriza los **sobrepasos** (SFAROVR)?"),
        dict(id="U24", titulo="Curso especial para egresados"),
    ]),
    ("Docentes", [
        dict(id="U26", titulo="Qué hace Registros Académicos en la carga lectiva"),
        dict(id="U28", titulo="Carga por contrato o por regla de carga",
             texto="Los contratos `ES` y `FC` ya están creados. (a) ¿La carga se mide **por contrato** (SIAFLCT y SIACONA, "
                   "lo recomendado en ADR-020) o **por regla de carga** (SIAFLRT y SIAASGN)? (b) ¿Quién crea en STVCNTR "
                   "los códigos de regla de cada contrato y en STVNIST los tipos de labor no educativa (matrícula, "
                   "coordinación, atención a estudiantes)?"),
        dict(id="U16", titulo="Separación de un docente"),
    ]),
    ("Servicios y seguimiento", [
        dict(id="U14", titulo="Área que recibe las quejas contra docentes"),
        dict(id="U15", titulo="Tipos de comentario de seguimiento"),
        dict(id="U07", titulo="¿La tutoría aplica a los centros?"),
    ]),
    ("Validación contra SEUSS", [
        dict(id="U11", titulo="Totales por centro y fecha de corte"),
    ]),
]

NECESITO = [
    ("Datos para completar las pruebas", [
        dict(id="U33", titulo="Alumno de prueba de Emprendimiento"),
        dict(id="U34", titulo="Datos que faltan de los NRC de prueba"),
    ]),
]

ESTADOS = {"validado": "Validado", "por probar": "Por probar", "pendiente": "Pendiente",
           "no aplica": "No aplica", "por definir": "Por definir"}
COLOR_ESTADO = {"validado": "E3F4DA", "por probar": "FFF4D6", "pendiente": "FFE8D9",
                "no aplica": "EEEEEE", "por definir": "EDE4F7"}


# ---------------------------------------------------------------------------------------------
def leer_registro():
    """{ID: texto de la duda} de las tablas de dudas abiertas del registro."""
    dudas = {}
    for linea in REGISTRO.read_text(encoding="utf-8").splitlines():
        m = re.match(r"^\| ([EU]\d{2}) \| (.+?) \| [^|]+ \| [^|]+ \|$", linea)
        if m:
            dudas[m.group(1)] = m.group(2).strip()
    return dudas


def significados():
    tabla = {}
    for k, v in {**common.PAGINAS, **common.SIGLAS}.items():
        tabla[k] = v[0] if isinstance(v, tuple) else v
    tabla.update(EXTRA)
    return tabla


class Escritor:
    def __init__(self):
        self.doc = Document()
        self.sig = significados()
        self.explicados = set()   # códigos ya explicados (la primera vez llevan su significado)
        self.usados = []          # para el glosario
        self._estilos()
        self._pagina()

    # ---- estilos y página ------------------------------------------------------------------
    def _estilos(self):
        st = self.doc.styles
        normal = st["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(10.5)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        normal.paragraph_format.space_after = Pt(4)
        normal.paragraph_format.line_spacing = 1.1
        for nombre, tam, color, antes in (("Title", 24, MORADO_OSC, 0), ("Heading 1", 16, MORADO, 14),
                                          ("Heading 2", 13, MORADO_OSC, 12), ("Heading 3", 11.5, "2B2140", 10)):
            s = st[nombre]
            s.font.name = "Calibri"
            s.font.size = Pt(tam)
            s.font.bold = True
            s.font.color.rgb = RGBColor.from_string(color)
            # Sin las fuentes del tema, para que Word use Calibri en los títulos.
            rfonts = s.element.rPr.find(qn("w:rFonts")) if s.element.rPr is not None else None
            if rfonts is not None:
                for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                    rfonts.attrib.pop(qn(a), None)
            s.paragraph_format.space_before = Pt(antes)
            s.paragraph_format.space_after = Pt(4)
            s.paragraph_format.keep_with_next = True

    def _pagina(self):
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2)
        sec.top_margin, sec.bottom_margin = Cm(2.6), Cm(2)
        sec.header_distance, sec.footer_distance = Cm(1), Cm(1)
        img = SKILL / "scripts" / "pdf" / "img"
        cab = sec.header.paragraphs[0]
        _solo_tab_derecha(self.doc, cab, "Header")
        cab.add_run().add_picture(str(img / "uss.png"), height=Cm(0.9))
        cab.add_run("\t")
        cab.add_run().add_picture(str(img / "ellucian.png"), height=Cm(0.7))
        _borde_inferior(cab, MORADO)
        pie = sec.footer.paragraphs[0]
        _solo_tab_derecha(self.doc, pie, "Footer")
        r = pie.add_run(f"Elaborado por: {AUTOR} · Centros Empresariales USS")
        r.font.size, r.font.color.rgb = Pt(8.5), RGBColor.from_string(GRIS)
        r = pie.add_run("\tPágina ")
        r.font.size, r.font.color.rgb = Pt(8.5), RGBColor.from_string(GRIS)
        _campo(pie, "PAGE")
        r = pie.add_run(" de ")
        r.font.size, r.font.color.rgb = Pt(8.5), RGBColor.from_string(GRIS)
        _campo(pie, "NUMPAGES")

    # ---- texto con **negrita**, `código` y siglas explicadas -----------------------------------
    def texto(self, parrafo, md, tam=None, color=None, cursiva=False, explicar=True):
        for trozo in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", md):
            if not trozo:
                continue
            negrita = trozo.startswith("**")
            codigo = trozo.startswith("`")
            limpio = trozo.strip("*`")
            partes = self._explicar(limpio) if explicar else [(limpio, False)]
            for t, es_sig in partes:
                r = parrafo.add_run(t)
                r.bold = negrita
                r.italic = cursiva or es_sig
                if codigo and not es_sig:
                    r.font.name = "Consolas"
                    r.font.size = Pt((tam or 10.5) - 1)
                elif tam:
                    r.font.size = Pt(tam)
                if es_sig:
                    r.font.color.rgb = RGBColor.from_string(GRIS)
                elif color:
                    r.font.color.rgb = RGBColor.from_string(color)
        return parrafo

    def _explicar(self, t):
        """Agrega «(significado)» después de la primera aparición de cada código del glosario."""
        salida, pos = [], 0
        for m in re.finditer(r"\b[A-Z][A-Z0-9_]{1,9}\b", t):
            cod = m.group(0)
            if cod in self.sig and cod not in self.explicados:
                self.explicados.add(cod)
                self.usados.append(cod)
                salida.append((t[pos:m.end()], False))
                salida.append((f" ({self.sig[cod]})", True))
                pos = m.end()
        salida.append((t[pos:], False))
        return [s for s in salida if s[0]]

    # ---- piezas ------------------------------------------------------------------------------
    def portada(self, total_e, total_u, total_n):
        d = self.doc
        p = d.add_paragraph(style="Title")
        p.add_run("Preguntas para Ellucian")
        p.paragraph_format.space_after = Pt(2)
        p = d.add_paragraph()
        r = p.add_run(f"Centros Empresariales USS · Reunión del {FECHA_REUNION}")
        r.font.size, r.bold, r.font.color.rgb = Pt(13), True, RGBColor.from_string(VERDE)
        p = d.add_paragraph()
        r = p.add_run(f"Idiomas (Inglés) · Computación (Informática) · Emprendimiento    |    Elaborado por: {AUTOR}")
        r.font.size, r.font.color.rgb = Pt(9.5), RGBColor.from_string(GRIS)
        _borde_inferior(p, MORADO)

        caja = self._caja("Cómo usar este documento", "F5F0FB", MORADO)
        pasos = [
            "Revise cada pregunta con Ellucian; el código (E01, U32…) es el mismo del registro de dudas.",
            "Escriba la respuesta en el recuadro «Respuesta», con el nombre de quien respondió.",
            "Marque el estado: «Resuelta» o «Sigue abierta» (y quién la verá).",
            "Al terminar, envíe este Word al asistente: pasa cada respuesta al registro (C##), actualiza los datos de "
            "cada centro y el asistente aprende la respuesta.",
        ]
        for i, t in enumerate(pasos, 1):
            q = caja.add_paragraph()
            q.paragraph_format.space_after = Pt(2)
            r = q.add_run(f"{i}.  ")
            r.bold, r.font.color.rgb = True, RGBColor.from_string(MORADO)
            self.texto(q, t, explicar=False)

        p = d.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        r = p.add_run("Contenido")
        r.bold, r.font.size, r.font.color.rgb = True, Pt(12), RGBColor.from_string(MORADO_OSC)
        for titulo, n, bloques in (("1. Para Ellucian", total_e, ELLUCIAN), ("2. Para la USS (preparar con los jefes)", total_u, USS),
                                   ("3. Lo que necesito para seguir aprendiendo", total_n, NECESITO)):
            q = d.add_paragraph()
            q.paragraph_format.space_after = Pt(1)
            r = q.add_run(f"{titulo}")
            r.bold = True
            r = q.add_run(f"  ·  {n} preguntas" if n else "")
            r.font.color.rgb = RGBColor.from_string(GRIS)
            for nombre, items in bloques:
                q = d.add_paragraph()
                q.paragraph_format.left_indent = Cm(0.6)
                q.paragraph_format.space_after = Pt(0)
                r = q.add_run(f"{nombre}: " )
                r.font.size = Pt(9.5)
                r = q.add_run(", ".join(it["id"] for it in items))
                r.font.size, r.font.color.rgb = Pt(9.5), RGBColor.from_string(GRIS)
        q = d.add_paragraph()
        q.paragraph_format.left_indent = Cm(0.6)
        r = q.add_run("Pruebas pendientes en TEST, por script y centro · Cómo enviarme lo que falta · Glosario")
        r.font.size, r.font.color.rgb = Pt(9.5), RGBColor.from_string(GRIS)

    def _caja(self, titulo, fondo, borde):
        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        celda = t.cell(0, 0)
        _bordes(t, borde, 8)
        _sombra(celda, fondo)
        p = celda.paragraphs[0]
        r = p.add_run(titulo)
        r.bold, r.font.color.rgb = True, RGBColor.from_string(MORADO_OSC)
        return celda

    def parte(self, titulo, intro):
        p = self.doc.add_paragraph()
        p.add_run().add_break(WD_BREAK.PAGE)
        p.paragraph_format.space_after = Pt(0)
        self.doc.add_heading(titulo, level=1)
        q = self.doc.add_paragraph()
        self.texto(q, intro, color=GRIS, explicar=False)

    def pregunta(self, item, dudas):
        d = self.doc
        texto = item.get("texto") or dudas.get(item["id"])
        if not texto:
            raise SystemExit(f"{item['id']} no está entre las dudas abiertas del registro")
        texto = re.sub(r"^\*\*[^*]+[:?]\*\*\s*", "", texto)
        texto = re.sub(r"^([¿¡«(]*)(\w)", lambda m: m.group(1) + m.group(2).upper(), texto)
        h = d.add_heading(level=3)
        r = h.add_run(f"{item['id']}  ")
        r.font.color.rgb = RGBColor.from_string(MORADO)
        h.add_run(item["titulo"])
        p = d.add_paragraph()
        p.paragraph_format.keep_with_next = True
        self.texto(p, texto)
        if item.get("sabemos"):
            p = d.add_paragraph()
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.left_indent = Cm(0.3)
            r = p.add_run("Lo que ya sabemos: ")
            r.bold, r.font.size, r.font.color.rgb = True, Pt(9.5), RGBColor.from_string(VERDE)
            self.texto(p, item["sabemos"], tam=9.5, color="4A4458")
        self.recuadro()

    def recuadro(self):
        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        _bordes(t, "B9A3D9", 6)
        fila = t.rows[0]
        _no_partir(fila)
        celda = t.cell(0, 0)
        p = celda.paragraphs[0]
        r = p.add_run("RESPUESTA")
        r.bold, r.font.size, r.font.color.rgb = True, Pt(8), RGBColor.from_string(MORADO)
        for _ in range(3):
            celda.add_paragraph()
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(8)
        r = p.add_run("Respondió: ____________________________     Estado:   ☐ Resuelta    ☐ Sigue abierta    ☐ Lo revisa Ellucian")
        r.font.size, r.font.color.rgb = Pt(8.5), RGBColor.from_string(GRIS)

    def bloques(self, bloques, dudas):
        for nombre, items in bloques:
            self.doc.add_heading(nombre, level=2)
            for it in items:
                self.pregunta(it, dudas)

    def tabla_pruebas(self):
        com = json.loads((RAIZ / "centros" / "comun" / "datos.json").read_text(encoding="utf-8"))
        centros = [(c, json.loads((RAIZ / "centros" / c / "datos.json").read_text(encoding="utf-8")))
                   for c in ("idiomas", "computacion", "emprendimiento")]
        d = self.doc
        d.add_heading("Pruebas pendientes en TEST", level=2)
        p = d.add_paragraph()
        self.texto(p, "Estado de cada script en `datos.json`. En la última columna anote la fecha, el NRC, el alumno y el "
                      "resultado de lo que pruebe; con eso actualizo el script a «Validado».", color=GRIS, explicar=False)
        anchos = [Cm(1.1), Cm(5.0), Cm(2.2), Cm(2.4), Cm(2.6), Cm(3.7)]
        t = d.add_table(rows=1, cols=6)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        _bordes(t, "D6CCE6", 4)
        for i, (celda, txt) in enumerate(zip(t.rows[0].cells, ["N.º", "Script", "Idiomas", "Computación", "Emprendimiento", "Fecha y resultado"])):
            celda.width = anchos[i]
            _sombra(celda, MORADO)
            r = celda.paragraphs[0].add_run(txt)
            r.bold, r.font.size, r.font.color.rgb = True, Pt(8.5), RGBColor.from_string("FFFFFF")
        _repetir_encabezado(t.rows[0])
        for s in com["scripts"]:
            fila = t.add_row()
            _no_partir(fila)
            celdas = fila.cells
            valores = [s["id"], s["nombre"]]
            for c, datos in centros:
                valores.append(datos["scripts"].get(s["id"], {}).get("estado", ""))
            valores.append("")
            for i, (celda, v) in enumerate(zip(celdas, valores)):
                celda.width = anchos[i]
                p = celda.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                texto = ESTADOS.get(v, v) if 2 <= i <= 4 else v
                r = p.add_run(texto)
                r.font.size = Pt(8.5)
                if 2 <= i <= 4 and v in COLOR_ESTADO:
                    _sombra(celda, COLOR_ESTADO[v])
                    r.bold = v != "validado"
        _anchos(t, anchos)

    def como_enviar(self):
        d = self.doc
        d.add_heading("Cómo enviarme lo que falta", level=2)
        lineas = [
            "**Respuestas:** escríbalas en cada recuadro y envíeme este mismo Word. Leo cada código (E01, U32…) y paso la "
            "respuesta al registro.",
            "**Pruebas en TEST:** una captura de la pantalla final (con el periodo, el NRC y el alumno visibles) y, si sale un "
            "error, el mensaje completo.",
            "**Datos de un centro** (programa, curso, NRC, escala, parte de periodo): con el código exacto de Banner; si es un "
            "ejemplo, indíquelo.",
            "**Lo que diga Ellucian:** quién lo dijo y en qué reunión o chat; si contradice un instructivo, lo anoto.",
            "**Sin datos reales de alumnos:** ni correos, ni DNI, ni contraseñas (ADR-016).",
        ]
        for t in lineas:
            p = d.add_paragraph(style="List Bullet")
            self.texto(p, t, explicar=False)

    def glosario(self):
        d = self.doc
        d.add_heading("Glosario de páginas y siglas", level=2)
        t = d.add_table(rows=0, cols=2)
        _bordes(t, "E2DAF0", 4)
        for cod in sorted(self.usados):
            fila = t.add_row()
            a, b = fila.cells
            a.width, b.width = Cm(3.2), Cm(13.8)
            r = a.paragraphs[0].add_run(cod)
            r.bold, r.font.size, r.font.color.rgb = True, Pt(9), RGBColor.from_string(MORADO_OSC)
            r = b.paragraphs[0].add_run(self.sig[cod][:1].upper() + self.sig[cod][1:])
            r.font.size = Pt(9)
            for c in (a, b):
                c.paragraphs[0].paragraph_format.space_after = Pt(0)
        _anchos(t, [Cm(3.2), Cm(13.8)])

    def propiedades(self):
        zoom = self.doc.settings.element.find(qn("w:zoom"))
        if zoom is not None and zoom.get(qn("w:percent")) is None:
            zoom.set(qn("w:percent"), "100")
        cp = self.doc.core_properties
        cp.author = AUTOR
        cp.title = f"Preguntas para Ellucian · Reunión del {FECHA_REUNION}"
        cp.subject = "Centros Empresariales USS · Ellucian Banner"
        cp.keywords = "Ellucian, Banner, Centros Empresariales, dudas"


# ---- utilidades de XML --------------------------------------------------------------------------
def _campo(parrafo, instr):
    for tipo, texto in (("begin", None), (None, instr), ("separate", None), (None, "1"), ("end", None)):
        r = parrafo.add_run()
        r.font.size, r.font.color.rgb = Pt(8.5), RGBColor.from_string(GRIS)
        if tipo:
            e = OxmlElement("w:fldChar")
            e.set(qn("w:fldCharType"), tipo)
        elif texto == instr:
            e = OxmlElement("w:instrText")
            e.set(qn("xml:space"), "preserve")
            e.text = f" {instr} "
        else:
            e = OxmlElement("w:t")
            e.text = texto
        r._r.append(e)


def _borde_inferior(parrafo, color):
    ppr = parrafo._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "8"), ("w:space", "4"), ("w:color", color)):
        b.set(qn(k), v)
    bdr.append(b)
    ppr.insert_element_before(bdr, "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
                              "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi",
                              "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing",
                              "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
                              "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr",
                              "w:pPrChange")


def _bordes(tabla, color, tam):
    tblpr = tabla._tbl.tblPr
    bordes = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{lado}")
        for k, v in (("w:val", "single"), ("w:sz", str(tam)), ("w:space", "0"), ("w:color", color)):
            e.set(qn(k), v)
        bordes.append(e)
    tblpr.insert_element_before(bordes, "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption",
                                "w:tblDescription", "w:tblPrChange")
    mar = OxmlElement("w:tblCellMar")
    for lado, v in (("top", 60), ("left", 110), ("bottom", 60), ("right", 110)):
        e = OxmlElement(f"w:{lado}")
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tblpr.insert_element_before(mar, "w:tblLook", "w:tblCaption", "w:tblDescription", "w:tblPrChange")


def _anchos(tabla, anchos):
    tabla.autofit = False
    tblpr = tabla._tbl.tblPr
    w = tblpr.find(qn("w:tblW"))
    if w is None:
        w = OxmlElement("w:tblW")
        tblpr.insert_element_before(w, "w:jc", "w:tblCellSpacing", "w:tblInd", "w:tblBorders", "w:shd",
                                    "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription")
    w.set(qn("w:w"), str(sum(int(a.twips) for a in anchos)))
    w.set(qn("w:type"), "dxa")
    for col, ancho in zip(tabla._tbl.tblGrid.findall(qn("w:gridCol")), anchos):
        col.set(qn("w:w"), str(int(ancho.twips)))
    for fila in tabla.rows:
        for celda, ancho in zip(fila.cells, anchos):
            celda.width = ancho


def _solo_tab_derecha(doc, parrafo, estilo):
    """Quita las tabulaciones del estilo (Header/Footer traen una centrada) y deja una a la derecha."""
    tabs = parrafo.paragraph_format.tab_stops
    for t in doc.styles[estilo].paragraph_format.tab_stops:
        tabs.add_tab_stop(t.position, WD_TAB_ALIGNMENT.CLEAR)
    tabs.add_tab_stop(Cm(17), WD_TAB_ALIGNMENT.RIGHT)


def _sombra(celda, color):
    tcpr = celda._tc.get_or_add_tcPr()
    s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear")
    s.set(qn("w:color"), "auto")
    s.set(qn("w:fill"), color)
    tcpr.append(s)


def _no_partir(fila):
    trpr = fila._tr.get_or_add_trPr()
    e = OxmlElement("w:cantSplit")
    trpr.append(e)


def _repetir_encabezado(fila):
    trpr = fila._tr.get_or_add_trPr()
    e = OxmlElement("w:tblHeader")
    trpr.append(e)


def main():
    salida = Path(sys.argv[1] if len(sys.argv) > 1 else
                  RAIZ / "docs" / "dudas" / "PREGUNTAS PARA ELLUCIAN - REUNIÓN 27-10 - CENTROS EMPRESARIALES.docx")
    dudas = leer_registro()
    total = lambda bloques: sum(len(items) for _, items in bloques)
    w = Escritor()
    w.portada(total(ELLUCIAN), total(USS), total(NECESITO))
    w.parte("1. Para Ellucian",
            "Ordenadas por impacto: primero lo que hoy bloquea la configuración o el cierre, después los detalles.")
    w.bloques(ELLUCIAN, dudas)
    w.parte("2. Para la USS (preparar con los jefes)",
            "No son para Ellucian, pero sus respuestas definen datos de Banner. Conviene tenerlas a la mano en la reunión.")
    w.bloques(USS, dudas)
    w.parte("3. Lo que necesito para seguir aprendiendo",
            "Con estas respuestas y pruebas completo los datos de cada centro, los diagramas y el asistente.")
    w.bloques(NECESITO, dudas)
    w.tabla_pruebas()
    w.como_enviar()
    w.glosario()
    w.propiedades()
    salida.parent.mkdir(parents=True, exist_ok=True)
    w.doc.save(salida)
    print(f"{salida} · {total(ELLUCIAN)} para Ellucian, {total(USS)} para la USS, {total(NECESITO)} de datos")


if __name__ == "__main__":
    main()
