"""Lee las respuestas escritas en el Word de preguntas (build_preguntas.py) y las lista por código.

    python3 leer_respuestas.py "docs/dudas/PREGUNTAS PARA ELLUCIAN - REUNIÓN 27-10 - CENTROS EMPRESARIALES.docx" [--json]

Cada pregunta es un título «E01  …» seguido de su recuadro «RESPUESTA» y de la línea «Respondió: … Estado: …».
Un estado se toma como marcado si su casilla ☐ se cambió por ☒, ☑, X o se borró el resto de opciones.
"""
import json
import re
import sys

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

ESTADOS = ("Resuelta", "Sigue abierta", "Lo revisa Ellucian")


def bloques(doc):
    """Recorre el cuerpo en orden (párrafos y tablas)."""
    for hijo in doc.element.body.iterchildren():
        if hijo.tag.endswith("}p"):
            yield Paragraph(hijo, doc)
        elif hijo.tag.endswith("}tbl"):
            yield Table(hijo, doc)


def leer(ruta):
    doc = Document(ruta)
    salida, actual, esperando = [], None, False
    for b in bloques(doc):
        if isinstance(b, Paragraph):
            m = re.match(r"^([EU]\d{2})\s+(.*)$", b.text.strip())
            if m and b.style.name.startswith("Heading"):
                actual = {"id": m.group(1), "titulo": m.group(2), "respuesta": "", "respondio": "", "estado": ""}
                salida.append(actual)
                esperando = True
            elif b.style.name.startswith("Heading"):
                actual, esperando = None, False
            elif actual and b.text.strip().startswith("Respondió:"):
                linea = b.text
                quien = re.search(r"Respondió:\s*(.*?)\s*Estado:", linea)
                actual["respondio"] = (quien.group(1) if quien else "").strip(" _")
                marcados = [e for e in ESTADOS if re.search(r"[☒☑✔✓xX]\s*" + re.escape(e), linea)]
                if not marcados:
                    presentes = [e for e in ESTADOS if e in linea]
                    marcados = presentes if len(presentes) == 1 else []
                actual["estado"] = ", ".join(marcados)
        elif actual and esperando:
            # El recuadro que sigue al título de la pregunta.
            texto = "\n".join(p.text for p in b.cell(0, 0).paragraphs).strip()
            actual["respuesta"] = re.sub(r"^RESPUESTA\s*", "", texto).strip()
            esperando = False
    return salida


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    filas = leer(sys.argv[1])
    if "--json" in sys.argv:
        print(json.dumps(filas, ensure_ascii=False, indent=1))
        return
    con = [f for f in filas if f["respuesta"]]
    print(f"{len(con)} de {len(filas)} preguntas con respuesta\n")
    for f in con:
        print(f"{f['id']} · {f['titulo']}\n  Respuesta: {f['respuesta']}\n  Respondió: {f['respondio'] or '—'} · Estado: {f['estado'] or '—'}\n")


if __name__ == "__main__":
    main()
